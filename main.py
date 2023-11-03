import os
import platform
from PIL import Image
import PyPDF2
from docx import Document
from openpyxl import load_workbook
import argparse
if platform.system().lower() == "windows":
    from win32com import client

def clear_metadata_image(image_path):
    try:
        image = Image.open(image_path)
        data = list(image.getdata())
        image_without_exif = Image.new(image.mode, image.size)
        image_without_exif.putdata(data)
        image_without_exif.save(image_path)
    except Exception as e:
        print(f"Error Image: {e}")

def clear_metadata_pdf(pdf_path):
    try:
        reader = PyPDF2.PdfFileReader(pdf_path)
        writer = PyPDF2.PdfFileWriter()

        for i in range(reader.getNumPages()):
            writer.addPage(reader.getPage(i))

        writer.encrypt('') 

        with open(pdf_path, 'wb') as new_file:
            writer.write(new_file)
    except Exception as e:
        print(f"Error PDF: {e}")

def clear_metadata_word(doc_path):
    try:
        doc = Document(doc_path)
        doc.core_properties.author = ""
        doc.core_properties.title = ""
        doc.core_properties.subject = ""
        doc.core_properties.creator = ""
        doc.core_properties.keywords = ""
        doc.core_properties.last_modified_by = ""
        doc.core_properties.language = ""
        doc.core_properties.category = ""
        doc.core_properties.company = ""
        doc.core_properties.revision = 1
        doc.core_properties.program_name = ""
        doc.core_properties.content_status = ""

        doc.save(doc_path)
    except Exception as e:
        print(f"Error Docx")

def clear_metadata_doc(doc_path):
    try:
        word = client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(doc_path)
        props = ["Title","Subject","Author","Keywords","Comments","Template","Last Author","Application Name","Category","Format","Manager","Company","Hyperlink Base","Content status","Content type"]
        for p in props:
            doc.BuiltInDocumentProperties(p).Value = ""
        doc.BuiltInDocumentProperties("Revision").Value = 1
        # TODO: Last saved by

        doc.Save()
    except Exception as e:
        print(f"Error Doc: {e}")

    finally:
        doc.Close()
        word.Quit()

def clear_metadata_excel(excel_path):
    try:
        wb = load_workbook(excel_path)
        props = wb.properties
        props.creator = ""
        props.last_modified_by = ""
        props.title = ""
        props.subject = ""
        props.category = ""
        props.keywords = ""
        props.program = ""
        props.comments = ""
        props.company = ""
        props.revision = 0
        wb.save(excel_path)
    except Exception as e:
        print(f"Error 5")

def clear_metadata_folder(folder_path):
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            if file.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp')):
                clear_metadata_image(file_path)
            elif file.lower().endswith('.pdf'):
                clear_metadata_pdf(file_path)
            elif file.lower().endswith('.docx'):
                clear_metadata_word(file_path)
            elif file.lower().endswith('doc') and platform.system().lower() == "windows":
                clear_metadata_doc(file_path)
            elif file.lower().endswith(('.xlsx', '.xls')):
                clear_metadata_excel(file_path)
            else:
                print(f"Unsupported: {file}")


def main():
    parser = argparse.ArgumentParser(description="Meta Data Cleaner")
    parser.add_argument("folder_path",type=str, help="Clean This Folder", default="/test")
    args = parser.parse_args()

    if os.path.isdir(args.folder_path):
        clear_metadata_folder(args.folder_path)
    else:
        print("Folder Not Found")

    print("DONE")

if __name__ == "__main__":
    main()
